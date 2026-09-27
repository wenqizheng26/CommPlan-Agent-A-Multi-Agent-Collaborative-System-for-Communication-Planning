import hashlib
import importlib.util
import json
from pathlib import Path
import socket
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from planning.retrieval import convert as conversion
from planning.retrieval.convert import converted, convert
from planning.retrieval.documents import DocumentStore, cached_chunks

HAS_MARKITDOWN = importlib.util.find_spec('markitdown') is not None
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def pdf(texts):
    """A small text PDF, one page per text (Helvetica, ASCII)."""
    objects = ['<< /Type /Catalog /Pages 2 0 R >>',
               '<< /Type /Pages /Kids [%s] /Count %d >>' % (' '.join(f'{4 + 2*i} 0 R' for i in range(len(texts))), len(texts)),
               '<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
    for i, text in enumerate(texts):
        stream = f'BT /F1 12 Tf 72 720 Td ({text}) Tj ET'
        objects.append(f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] '
                       f'/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2*i} 0 R >>')
        objects.append(f'<< /Length {len(stream)} >>\nstream\n{stream}\nendstream')
    out, offsets = b'%PDF-1.4\n', []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f'{number} 0 obj\n{body}\nendobj\n'.encode('latin-1')
    start = len(out)
    out += f'xref\n0 {len(objects) + 1}\n0000000000 65535 f \n'.encode()
    out += b''.join(f'{o:010d} 00000 n \n'.encode() for o in offsets)
    return out + f'trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n'.encode()


def docx(path, paragraphs):
    style = lambda name: f'<w:pPr><w:pStyle w:val="{name}"/></w:pPr>' if name else ''
    body = ''.join(f'<w:p>{style(name)}<w:r><w:t>{text}</w:t></w:r></w:p>' for name, text in paragraphs)
    rels = 'http://schemas.openxmlformats.org/package/2006/relationships'
    kind = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('[Content_Types].xml', '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/></Types>')
        archive.writestr('_rels/.rels', f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="{rels}">'
            f'<Relationship Id="rId1" Type="{kind}/officeDocument" Target="word/document.xml"/></Relationships>')
        archive.writestr('word/_rels/document.xml.rels', f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="{rels}">'
            f'<Relationship Id="rId1" Type="{kind}/styles" Target="styles.xml"/></Relationships>')
        archive.writestr('word/styles.xml', f'<?xml version="1.0" encoding="UTF-8"?><w:styles xmlns:w="{W}">'
            '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/></w:style>'
            '<w:style w:type="paragraph" w:styleId="Heading2"><w:name w:val="heading 2"/></w:style></w:styles>')
        archive.writestr('word/document.xml', f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="{W}"><w:body>{body}</w:body></w:document>')


class ConversionTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.added = self.root / 'knowledge/sources/added'
        self.added.mkdir(parents=True)
        (self.root / 'knowledge/documents').mkdir(parents=True)
        self.records = []

    def register(self, name, data):
        path = self.added / name
        path.write_bytes(data) if isinstance(data, bytes) else data(path)
        record = dict(doc_id=f'added-{len(self.records) + 1}', title=name, version='1', language='zh',
                      source='本机添加：' + name, local_path='knowledge/sources/added/' + name,
                      sha256=hashlib.sha256(path.read_bytes()).hexdigest(), simulated=False, redistributable=False)
        self.records.append(record)
        (self.root / 'knowledge/documents/manifest.json').write_text(
            json.dumps(dict(schema_version=1, documents=self.records), ensure_ascii=False), encoding='utf-8')
        return record

    def test_pdf_pages_are_locators_with_or_without_markitdown(self):
        self.register('link.pdf', pdf(['Free space loss 92.4 dB', 'Received power -60 dBm']))
        with patch.object(socket.socket, 'connect', side_effect=AssertionError('unexpected network')):
            store = DocumentStore(self.root)
        self.assertEqual([(c['id'], c['sources'][0]['locator']) for c in store.chunks],
                         [('doc:added-1:p1-1', 'page 1, chunk 1'), ('doc:added-1:p2-1', 'page 2, chunk 1')])
        self.assertIn('92.4 dB', store.chunks[0]['description'])
        with patch.object(conversion, 'converter_for', side_effect=ImportError('markitdown')):
            self.assertEqual(convert(self.added / 'link.pdf'), converted(self.root, self.records[0]))

    @unittest.skipUnless(HAS_MARKITDOWN, 'MarkItDown is optional (requirements-docs.txt)')
    def test_office_and_web_formats_become_located_sections(self):
        from openpyxl import Workbook
        from pptx import Presentation

        def slides(path):
            deck = Presentation()
            for title, text in (('链路预算', '发射功率 40 dBm'), ('门限', '接收灵敏度 -97 dBm')):
                slide = deck.slides.add_slide(deck.slide_layouts[1])
                slide.shapes.title.text, slide.placeholders[1].text = title, text
            deck.save(path)

        def sheet(path):
            book = Workbook()
            book.active.title = '设备'
            book.active.append(['型号', '发射功率 dBm'])
            book.active.append(['XX-300', 40])
            book.save(path)

        self.register('manual.docx', lambda path: docx(path, [('Heading1', 'XX-300 手册'), ('Heading2', '规格'), (None, '发射功率 40 dBm')]))
        self.register('deck.pptx', slides)
        self.register('specs.xlsx', sheet)
        self.register('page.html', '<html><head><script>alert("x")</script></head><body><h1>站点</h1><p>A站 坐标</p></body></html>'.encode('utf-8'))
        self.register('notes.txt', '接收灵敏度 -97 dBm。'.encode('gbk'))
        with patch.object(socket.socket, 'connect', side_effect=AssertionError('unexpected network')):
            store = DocumentStore(self.root)
        self.assertEqual({s['status'] for s in store.status}, {'ready'})
        located = {c['sources'][0]['locator']: c['description'] for c in store.chunks}
        self.assertIn('发射功率 40 dBm', located['XX-300 手册 / 规格, chunk 1'])
        self.assertIn('接收灵敏度 -97 dBm', located['幻灯片 2 / 门限, chunk 1'])
        self.assertIn('| XX-300 | 40 |', located['设备, chunk 1'])
        self.assertNotIn('alert', ' '.join(located))
        self.assertEqual(located['正文, chunk 1'], '接收灵敏度 -97 dBm。')

    def test_cache_is_checked_and_replaced_when_the_source_changes(self):
        record = self.register('link.pdf', pdf(['Free space loss 92.4 dB']))
        text = converted(self.root, record)
        cache = self.root / f"knowledge/sources/converted/added-1.{record['sha256'][:16]}.md"
        self.assertTrue(cache.read_text(encoding='utf-8').endswith(text))
        cache.write_text(cache.read_text(encoding='utf-8').replace('92.4', '99.9'), encoding='utf-8')
        self.assertEqual(converted(self.root, record), text)  # an edited cache is converted again
        self.assertNotIn('99.9', cache.read_text(encoding='utf-8'))
        (self.added / 'link.pdf').write_bytes(pdf(['Free space loss 100.1 dB']))
        record = dict(record, sha256=hashlib.sha256((self.added / 'link.pdf').read_bytes()).hexdigest())
        self.assertIn('100.1', converted(self.root, record))
        self.assertEqual(len(list(cache.parent.glob('added-1.*.md'))), 1)

    def test_damaged_or_unknown_files_are_reported_without_blocking_others(self):
        self.register('good.pdf', pdf(['Free space loss 92.4 dB']))
        self.register('broken.docx', b'PK\x03\x04 not really a document')
        self.register('broken.pdf', b'%PDF-1.7\ninvalid')
        self.register('image.png', b'\x89PNG\r\n')
        cached_chunks.cache_clear()
        store = DocumentStore(self.root)
        status = {s['doc_id']: (s['status'], s.get('code')) for s in store.status}
        self.assertEqual(status['added-1'], ('ready', None))
        self.assertEqual(status['added-2'][0], 'unreadable')
        if HAS_MARKITDOWN:
            self.assertEqual(status['added-2'][1], 'DOCUMENT_UNREADABLE')
        self.assertEqual(status['added-3'], ('unreadable', 'DOCUMENT_PDF_UNREADABLE'))
        self.assertEqual(status['added-4'], ('unreadable', 'DOCUMENT_FORMAT'))
        self.assertTrue(all(c['doc_id'] == 'added-1' for c in store.chunks))


if __name__ == '__main__':
    unittest.main()
